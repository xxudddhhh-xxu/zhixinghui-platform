"""文件解析: txt/md 直接读取, pdf 用 pypdf, docx 用 python-docx。"""

import io
import logging

logger = logging.getLogger("rag.parser")

TEXT_EXT = (".txt", ".md", ".markdown", ".csv", ".log")


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(TEXT_EXT):
        return data.decode("utf-8", errors="ignore")
    if name.endswith(".pdf"):
        return _pdf(data)
    if name.endswith(".docx"):
        return _docx(data)
    raise ValueError(f"不支持的文件类型: {filename}(支持 txt/md/pdf/docx)")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n\n".join(pages)


def _docx(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)
